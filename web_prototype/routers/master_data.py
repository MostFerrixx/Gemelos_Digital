# -*- coding: utf-8 -*-
"""Endpoints de DATOS MAESTROS (Excel maestro, mapa TMX, tablas de la BD).

Por que existe este router
--------------------------
Los datos maestros del almacen (ubicaciones, catalogo de SKU, zonas de staging,
muelles) viven en `layouts/Warehouse_Logic.xlsx`, pero el motor NO los lee de
ahi: los lee de `warehouse.db`. El Excel solo se usa si la base no existe.

Consecuencia que este router viene a resolver: **editar el Excel no tenia
ningun efecto** hasta que alguien corriera `run_migration.py` por consola, y la
UI no avisaba de eso en ninguna parte. Los botones "Examinar" de la pestana
Layout y Datos, ademas, no hacian nada.

Flujo que se expone:
    subir  ->  validar (sin aplicar)  ->  aplicar (migrar a la BD)  ->  ver

Ver `docs/PLAN_DATOS_MAESTROS_WEB.md`.
"""
import os
import shutil
import sqlite3
import subprocess
import sys
import time
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel

from web_prototype.app_state import PROJECT_ROOT

router = APIRouter()

DB_PATH = os.path.join(PROJECT_ROOT, "warehouse.db")
UPLOADS_DIR = os.path.join(PROJECT_ROOT, "uploads")

# Hojas que el motor necesita, con sus columnas obligatorias.
# Espejo de lo que consume data_manager._load_excel_data / el importador.
HOJAS_REQUERIDAS: Dict[str, List[str]] = {
    "PickingLocations": ["x", "y", "pick_sequence", "WorkArea"],
}
HOJAS_OPCIONALES: Dict[str, List[str]] = {
    "SkuCatalog": ["sku_code", "volumen_m3", "peso_kg", "clase_manejo"],
    "OutboundStaging": ["staging_id", "x", "y"],
    "InboundDocks": ["dock_id", "x", "y"],
}

# Tablas de la BD que se pueden consultar, y la hoja de la que provienen.
TABLAS = {
    "locations": "PickingLocations",
    "sku_catalog": "SkuCatalog",
    "staging_areas": "OutboundStaging",
    "inbound_docks": "InboundDocks",
}


class AplicarRequest(BaseModel):
    """Aplica un Excel ya subido (o el vigente si no se pasa ninguno)."""
    excel_path: Optional[str] = None


def _conn():
    if not os.path.exists(DB_PATH):
        raise HTTPException(status_code=404,
                            detail="No existe warehouse.db. Subi y aplica el Excel maestro.")
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def _ruta_segura(rel: str) -> str:
    """Resuelve una ruta relativa al proyecto y bloquea path traversal."""
    resuelto = os.path.realpath(os.path.join(PROJECT_ROOT, rel))
    raiz = os.path.realpath(PROJECT_ROOT)
    if not resuelto.startswith(raiz + os.sep):
        raise HTTPException(status_code=400, detail="Ruta fuera del proyecto.")
    return resuelto


def _validar_excel(path: str, tmx_path: Optional[str] = None) -> Dict[str, Any]:
    """Abre el Excel y comprueba hojas, columnas y (QA-6.8) que sus
    coordenadas existan y sean transitables en el mapa configurado. NO aplica
    nada."""
    import openpyxl

    errores: List[str] = []
    avisos: List[str] = []
    resumen: Dict[str, int] = {}

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return {"valido": False, "errores": ["No se pudo abrir el archivo: %s" % e],
                "avisos": [], "resumen": {}}

    def _columnas(hoja):
        fila = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(c).strip() for c in fila if c is not None]

    for hoja, requeridas in HOJAS_REQUERIDAS.items():
        if hoja not in wb.sheetnames:
            errores.append("Falta la hoja obligatoria '%s'." % hoja)
            continue
        ws = wb[hoja]
        cols = _columnas(ws)
        faltantes = [c for c in requeridas if c not in cols]
        if faltantes:
            errores.append("La hoja '%s' no tiene la(s) columna(s): %s."
                           % (hoja, ", ".join(faltantes)))
        resumen[hoja] = max(ws.max_row - 1, 0)

    for hoja, requeridas in HOJAS_OPCIONALES.items():
        if hoja not in wb.sheetnames:
            avisos.append("Sin la hoja '%s' (opcional): se usaran los valores por defecto."
                          % hoja)
            resumen[hoja] = 0
            continue
        ws = wb[hoja]
        cols = _columnas(ws)
        faltantes = [c for c in requeridas if c not in cols]
        if faltantes:
            errores.append("La hoja '%s' no tiene la(s) columna(s): %s."
                           % (hoja, ", ".join(faltantes)))
        resumen[hoja] = max(ws.max_row - 1, 0)

    # QA-6.8: coordenadas contra el mapa (antes un carril sobre un rack o una
    # ubicacion fuera del mapa se aplicaba igual y la corrida se rompia).
    if not errores:
        if tmx_path is None:
            try:
                from web_prototype.app_state import config_manager
                tmx_path = config_manager.load_config().get("layout_file")
            except Exception:
                tmx_path = None
        if tmx_path and not os.path.isabs(tmx_path):
            tmx_path = os.path.join(PROJECT_ROOT, tmx_path)
        if tmx_path and os.path.exists(tmx_path):
            puntos = []
            for hoja, etiqueta, col_id in (("PickingLocations", "ubicacion en fila", None),
                                           ("OutboundStaging", "zona de salida", "staging_id"),
                                           ("InboundDocks", "muelle", "dock_id")):
                if hoja not in wb.sheetnames:
                    continue
                ws = wb[hoja]
                cols = _columnas(ws)
                ix, iy = cols.index("x"), cols.index("y")
                iid = cols.index(col_id) if col_id in cols else None
                for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        x, y = int(fila[ix]), int(fila[iy])
                    except (TypeError, ValueError, IndexError):
                        continue
                    ident = str(fila[iid]) if iid is not None else str(n)
                    puntos.append((etiqueta, ident, x, y))
            try:
                ancho, alto, bloqueadas = _leer_mapa(tmx_path)
                errores += _cruzar_puntos(puntos, ancho, alto, bloqueadas,
                                          "este Excel (mapa %s)" % os.path.basename(tmx_path))
            except Exception as e:
                avisos.append("No se pudo cruzar con el mapa %s: %s" % (tmx_path, e))
        else:
            avisos.append("No hay mapa configurado para verificar las coordenadas.")

    wb.close()
    return {"valido": not errores, "errores": errores, "avisos": avisos, "resumen": resumen}


@router.post("/api/master-data/upload-excel")
async def upload_excel(file: UploadFile = File(...)):
    """Sube el Excel maestro y lo VALIDA. No toca la base todavia."""
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos .xlsx")

    os.makedirs(UPLOADS_DIR, exist_ok=True)
    destino = os.path.join(UPLOADS_DIR, "master_%s" % os.path.basename(file.filename))
    with open(destino, "wb") as f:
        f.write(await file.read())

    reporte = _validar_excel(destino)
    reporte["excel_path"] = os.path.relpath(destino, PROJECT_ROOT)
    reporte["nombre"] = os.path.basename(file.filename)
    return reporte


def _celdas_bloqueadas(ruta: str, ancho: int, alto: int):
    """Misma regla que el motor (layout_manager): toda celda es transitable
    salvo que algun tile de una capa visible diga walkable=false."""
    import pytmx
    tmx = pytmx.TiledMap(ruta)
    bloqueadas = set()
    for idx, capa in enumerate(tmx.visible_layers):
        if not hasattr(capa, "data"):
            continue
        for y in range(alto):
            for x in range(ancho):
                try:
                    props = tmx.get_tile_properties(x, y, idx)
                except (AttributeError, IndexError, KeyError):
                    continue
                if not props:
                    continue
                w = props.get("walkable", "true")
                ok = w.lower() == "true" if isinstance(w, str) else bool(w)
                if not ok:
                    bloqueadas.add((x, y))
    return bloqueadas


def _n_puntos(lista) -> str:
    return "1 punto" if len(lista) == 1 else "%d puntos" % len(lista)


def _ejemplos(lista, n=5):
    return ", ".join(lista[:n]) + (" y %d mas" % (len(lista) - n) if len(lista) > n else "")


def _cruzar_puntos(puntos, ancho: int, alto: int, bloqueadas, origen: str) -> List[str]:
    """Errores si algun punto (etiqueta, id, x, y) cae fuera del mapa o sobre
    una celda bloqueada. `origen` dice de donde vienen los puntos."""
    errores = []
    fuera = ["%s %s (%d, %d)" % p for p in puntos
             if not (0 <= p[2] < ancho and 0 <= p[3] < alto)]
    tapados = ["%s %s (%d, %d)" % p for p in puntos
               if 0 <= p[2] < ancho and 0 <= p[3] < alto and (p[2], p[3]) in bloqueadas]
    if fuera:
        errores.append("%s de %s %s FUERA del mapa de %d x %d: %s."
                       % (_n_puntos(fuera), origen, "queda" if len(fuera) == 1 else "quedan",
                          ancho, alto, _ejemplos(fuera)))
    if tapados:
        errores.append("%s de %s %s sobre celdas bloqueadas (racks o paredes): %s. "
                       "Los operarios no podrian llegar."
                       % (_n_puntos(tapados), origen, "cae" if len(tapados) == 1 else "caen",
                          _ejemplos(tapados)))
    return errores


def _leer_mapa(ruta: str):
    """(ancho, alto, celdas bloqueadas) con la misma regla que el motor."""
    import xml.etree.ElementTree as ET
    raiz = ET.parse(ruta).getroot()
    ancho, alto = int(raiz.get("width", 0)), int(raiz.get("height", 0))
    return ancho, alto, _celdas_bloqueadas(ruta, ancho, alto)


def validar_tmx(ruta: str, db_path: Optional[str] = None) -> Dict[str, Any]:
    """QA-6.6: un mapa se acepta solo si el simulador puede usarlo CON los
    datos maestros en uso. Antes bastaba un <map> (se aceptaba un mapa vacio
    de 10x10 aunque las ubicaciones llegan a x=31, y=39)."""
    import xml.etree.ElementTree as ET
    errores: List[str] = []
    avisos: List[str] = []
    try:
        raiz = ET.parse(ruta).getroot()
    except Exception as e:
        return {"valido": False, "errores": ["No es un archivo TMX valido (XML ilegible: %s)." % e],
                "avisos": [], "resumen": {}}
    if raiz.tag != "map":
        return {"valido": False, "errores": ["No es un mapa de Tiled: falta el elemento <map>."],
                "avisos": [], "resumen": {}}
    try:
        ancho, alto = int(raiz.get("width", 0)), int(raiz.get("height", 0))
    except ValueError:
        ancho = alto = 0
    capas = [c.get("name") for c in raiz.findall("layer")]
    if ancho <= 0 or alto <= 0:
        errores.append("El mapa no tiene tamanio (ancho x alto = %d x %d)." % (ancho, alto))
    if not capas:
        errores.append("El mapa no tiene capas de tiles: el simulador no sabria donde hay "
                       "racks, pasillos ni muelles.")
    resumen = {"ancho": ancho, "alto": alto, "capas": len(capas)}
    if errores:
        return {"valido": False, "errores": errores, "avisos": avisos,
                "resumen": resumen, "capas": capas}

    try:
        bloqueadas = _celdas_bloqueadas(ruta, ancho, alto)
    except Exception as e:
        return {"valido": False, "errores": ["El mapa esta danado o incompleto y el simulador tampoco "
                                        "puede abrirlo (%s). Revisalo en Tiled." % e],
                "avisos": avisos, "resumen": resumen, "capas": capas}
    resumen["celdas_bloqueadas"] = len(bloqueadas)
    if not bloqueadas:
        avisos.append("Ninguna celda esta bloqueada: el mapa no tiene racks ni paredes.")

    # Cruce con los datos maestros que usa el motor (warehouse.db).
    if db_path and os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        try:
            puntos = []
            for sql, etiqueta in (
                    ("SELECT location_id, legacy_x, legacy_y FROM locations", "ubicacion"),
                    ("SELECT staging_id, legacy_x, legacy_y FROM staging_areas", "zona de salida"),
                    ("SELECT dock_id, x, y FROM inbound_docks", "muelle")):
                try:
                    puntos += [(etiqueta, str(i), int(x), int(y))
                               for i, x, y in conn.execute(sql).fetchall()
                               if x is not None and y is not None]
                except sqlite3.Error:
                    continue
        finally:
            conn.close()
        errores += _cruzar_puntos(puntos, ancho, alto, bloqueadas, "los datos maestros en uso")
        if errores:
            errores.append("Usa un mapa que los contenga o aplica un Excel hecho para este mapa.")
        resumen["puntos_verificados"] = len(puntos)
    return {"valido": not errores, "errores": errores, "avisos": avisos,
            "resumen": resumen, "capas": capas}


@router.post("/api/master-data/upload-tmx")
async def upload_tmx(file: UploadFile = File(...)):
    """Sube un mapa .tmx y comprueba que se pueda abrir."""
    if not file.filename.lower().endswith(".tmx"):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos .tmx")

    os.makedirs(UPLOADS_DIR, exist_ok=True)
    destino = os.path.join(UPLOADS_DIR, "map_%s" % os.path.basename(file.filename))
    with open(destino, "wb") as f:
        f.write(await file.read())

    info = validar_tmx(destino, DB_PATH)
    info["tmx_path"] = os.path.relpath(destino, PROJECT_ROOT)
    info["nombre"] = os.path.basename(file.filename)
    return info


@router.post("/api/master-data/apply")
def aplicar_excel(request: AplicarRequest):
    """Importa el Excel a `warehouse.db` (el origen que SI lee el motor).

    Hace BACKUP de la base antes de tocarla. La importacion RECONSTRUYE las
    tablas: cualquier edicion hecha desde la web se reemplaza por lo que traiga
    el Excel, y el stock inicial vuelve al del archivo (decision del Director,
    2026-09-08).
    """
    from web_prototype.app_state import config_manager

    if request.excel_path:
        excel = _ruta_segura(request.excel_path)
    else:
        cfg = config_manager.load_config()
        excel = _ruta_segura(cfg.get("sequence_file", "layouts/Warehouse_Logic.xlsx"))

    if not os.path.exists(excel):
        raise HTTPException(status_code=404, detail="No se encontro el Excel: %s" % excel)

    reporte = _validar_excel(excel)
    if not reporte["valido"]:
        return {"success": False, "errors": reporte["errores"]}

    backup = None
    if os.path.exists(DB_PATH):
        backup = DB_PATH + ".bak"
        shutil.copy2(DB_PATH, backup)

    proc = subprocess.run(
        [sys.executable, os.path.join(PROJECT_ROOT, "run_migration.py"), "--excel", excel],
        cwd=PROJECT_ROOT, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=600,
    )

    if proc.returncode != 0:
        # Restaurar la base: mejor quedarse como estaba que a medio migrar.
        if backup and os.path.exists(backup):
            shutil.copy2(backup, DB_PATH)
        cola = "\n".join((proc.stdout or "").splitlines()[-15:])
        return {"success": False,
                "errors": ["La importacion fallo (se restauro la base anterior)."],
                "log": cola}

    return {
        "success": True,
        "excel": os.path.relpath(excel, PROJECT_ROOT),
        "backup": os.path.relpath(backup, PROJECT_ROOT) if backup else None,
        "resumen": _resumen_bd(),
        "log": "\n".join((proc.stdout or "").splitlines()[-15:]),
    }


def _meta_aplicacion(conn) -> Dict[str, Any]:
    """BK-35: cuando y de que Excel se aplicaron los datos (si se sabe)."""
    try:
        meta = dict(conn.execute("SELECT clave, valor FROM master_data_meta").fetchall())
    except sqlite3.Error:
        return {}
    out = {}
    try:
        out["aplicado_en"] = float(meta["aplicado_en"])
    except (KeyError, TypeError, ValueError):
        pass
    if meta.get("excel"):
        out["excel_aplicado"] = os.path.relpath(meta["excel"], PROJECT_ROOT)
    return out


def _resumen_bd() -> Dict[str, Any]:
    if not os.path.exists(DB_PATH):
        return {"disponible": False}
    conn = sqlite3.connect(DB_PATH)
    try:
        datos = {}
        for tabla in TABLAS:
            try:
                datos[tabla] = conn.execute("SELECT COUNT(*) FROM %s" % tabla).fetchone()[0]
            except sqlite3.Error:
                datos[tabla] = None
        out = {"disponible": True, "conteos": datos,
               "actualizada": os.path.getmtime(DB_PATH)}
        out.update(_meta_aplicacion(conn))
        return out
    finally:
        conn.close()


@router.get("/api/master-data/summary")
def resumen():
    """Que datos maestros esta usando el motor, y si el Excel quedo mas nuevo.

    El aviso de desactualizacion es la parte importante: sin el, alguien edita
    el Excel, no ve ningun cambio en los resultados y no entiende por que.
    """
    from web_prototype.app_state import config_manager

    info = _resumen_bd()
    try:
        cfg = config_manager.load_config()
        excel = _ruta_segura(cfg.get("sequence_file", "layouts/Warehouse_Logic.xlsx"))
        info["excel"] = os.path.relpath(excel, PROJECT_ROOT)
        if os.path.exists(excel):
            info["excel_actualizado"] = os.path.getmtime(excel)
            info["excel_mas_nuevo"] = (
                info.get("disponible", False)
                and info["excel_actualizado"] > info.get("aplicado_en", info.get("actualizada", 0))
            )
        else:
            info["excel_existe"] = False
    except Exception as e:
        info["excel_error"] = str(e)
    return info


class Coordenada(BaseModel):
    """Una fila de las tablas chicas editables (zona de staging o muelle)."""
    id: int
    x: int
    y: int


class CoordenadasRequest(BaseModel):
    filas: List[Coordenada]


# Tablas chicas editables desde la web: (tabla, columna id, col x, col y).
# Solo estas: las grandes (ubicaciones, catalogo) se editan en el Excel, que es
# mejor herramienta para eso, y se traen con "Aplicar Excel".
EDITABLES = {
    "staging_areas": ("staging_id", "legacy_x", "legacy_y"),
    "inbound_docks": ("dock_id", "x", "y"),
}


def _mapa_para_validar():
    """(ancho, alto, es_transitable) del mapa TMX vigente, o (None, None, None).

    Se usa el LayoutManager REAL del motor en vez de reinterpretar el TMX aca:
    es la misma fuente de verdad que valida las coordenadas al correr. Si no se
    puede cargar (falta el archivo, falta una dependencia), se degrada a no
    validar geometria en vez de bloquear el guardado.
    """
    try:
        src_dir = os.path.join(PROJECT_ROOT, "src")
        if src_dir not in sys.path:
            sys.path.insert(0, src_dir)
        from web_prototype.app_state import config_manager
        from subsystems.simulation.layout_manager import LayoutManager

        tmx = config_manager.load_config().get("layout_file", "layouts/WH1.tmx")
        lm = LayoutManager(_ruta_segura(tmx), headless=True)
        return lm.grid_width, lm.grid_height, lm.is_walkable
    except Exception as e:
        print("[MASTER-DATA][WARN] No se pudo cargar el mapa para validar: %s" % e)
        return None, None, None


def _guardar_coordenadas(tabla: str, filas: List[Coordenada]) -> Dict[str, Any]:
    col_id, col_x, col_y = EDITABLES[tabla]

    ids = [f.id for f in filas]
    if len(set(ids)) != len(ids):
        raise HTTPException(status_code=400, detail="Hay identificadores repetidos.")
    for f in filas:
        if f.x < 0 or f.y < 0:
            raise HTTPException(status_code=400,
                                detail="Coordenadas negativas en el id %s." % f.id)

    # Contra el MAPA: el motor exige que estas coordenadas caigan dentro de la
    # grilla del TMX y aborta la corrida si no (DataManagerError). Sin este
    # chequeo, la web aceptaba un valor invalido y el error recien aparecia al
    # simular, lejos de donde se cometio.
    ancho, alto, transitable = _mapa_para_validar()
    if ancho is not None:
        fuera = [f for f in filas if not (0 <= f.x < ancho and 0 <= f.y < alto)]
        if fuera:
            raise HTTPException(
                status_code=400,
                detail=("Fuera del mapa (%d x %d): %s. Corregi las coordenadas."
                        % (ancho, alto,
                           ", ".join("id %s -> (%s, %s)" % (f.id, f.x, f.y) for f in fuera))))

        if transitable is not None:
            bloqueadas = []
            for f in filas:
                try:
                    if not transitable(f.x, f.y):
                        bloqueadas.append(f)
                except Exception:
                    pass  # ante la duda, no bloquear el guardado
            if bloqueadas:
                raise HTTPException(
                    status_code=400,
                    detail=("Estas celdas no son transitables en el mapa (hay un rack o "
                            "una pared): %s. Los agentes no podrian llegar."
                            % ", ".join("id %s -> (%s, %s)" % (f.id, f.x, f.y)
                                        for f in bloqueadas)))

    conn = _conn()
    try:
        existentes = {r[col_id] for r in conn.execute(
            "SELECT %s FROM %s" % (col_id, tabla))}
        desconocidos = [i for i in ids if i not in existentes]
        if desconocidos:
            raise HTTPException(
                status_code=400,
                detail=("Estos identificadores no existen: %s. Para agregar o quitar "
                        "filas, editalos en el Excel y usa 'Aplicar Excel'."
                        % ", ".join(str(d) for d in desconocidos)))

        for f in filas:
            conn.execute("UPDATE %s SET %s = ?, %s = ? WHERE %s = ?"
                         % (tabla, col_x, col_y, col_id), (f.x, f.y, f.id))
        conn.commit()
        return {"success": True, "actualizadas": len(filas)}
    finally:
        conn.close()


@router.put("/api/master-data/staging-areas")
def guardar_staging(request: CoordenadasRequest):
    """Actualiza las coordenadas de las zonas de salida.

    OJO: escribe en `warehouse.db`, no en el Excel. Un "Aplicar Excel" posterior
    reemplaza estos valores por los del archivo (decision del Director,
    2026-09-08: el Excel manda, avisando antes).
    """
    return _guardar_coordenadas("staging_areas", request.filas)


@router.put("/api/master-data/inbound-docks")
def guardar_docks(request: CoordenadasRequest):
    """Actualiza las coordenadas de los muelles de recepcion. Ver nota de arriba."""
    return _guardar_coordenadas("inbound_docks", request.filas)


def _leer_stock(limit: int, offset: int, q: str) -> Dict[str, Any]:
    """BK-35: stock INICIAL por ubicacion, el que usa cada corrida al arrancar.

    `inventory` es la copia de trabajo (cambia durante cada corrida); el stock
    con el que arranca la proxima es `inventory_baseline` si existe (H-38) o,
    si todavia no se corrio nada desde el ultimo "Aplicar Excel", `inventory`.
    """
    conn = _conn()
    try:
        hay_foto = conn.execute("SELECT count(*) FROM sqlite_master "
                                "WHERE name='inventory_baseline'").fetchone()[0] > 0
        stock = "COALESCE(b.qty_baseline, i.qty_available)" if hay_foto else "i.qty_available"
        union = "LEFT JOIN inventory_baseline b ON b.location_id = i.location_id" if hay_foto else ""
        base = ("FROM inventory i LEFT JOIN locations l ON l.location_id = i.location_id %s" % union)
        where, params = "", []
        if q:
            where = " WHERE i.location_id LIKE ? OR i.sku_code LIKE ? OR l.work_area LIKE ?"
            params = ["%%%s%%" % q] * 3
        total = conn.execute("SELECT COUNT(*) %s%s" % (base, where), params).fetchone()[0]
        filas = conn.execute(
            "SELECT i.location_id, i.sku_code, %s AS stock_inicial, l.work_area, "
            "l.legacy_x, l.legacy_y %s%s ORDER BY i.location_id LIMIT ? OFFSET ?"
            % (stock, base, where), params + [limit, offset]).fetchall()
        columnas = ["location_id", "sku_code", "stock_inicial", "work_area", "legacy_x", "legacy_y"]
        return {"tabla": "stock", "hoja": "PickingLocations (qty_initial)", "columnas": columnas,
                "total": total, "limit": limit, "offset": offset,
                "filas": [dict(zip(columnas, f)) for f in filas]}
    finally:
        conn.close()


@router.get("/api/master-data/table/{nombre}")
def leer_tabla(nombre: str, limit: int = 50, offset: int = 0, q: str = ""):
    """Contenido paginado de una tabla de datos maestros (solo lectura)."""
    if nombre == "stock":
        return _leer_stock(max(1, min(int(limit), 500)), max(0, int(offset)), q)
    if nombre not in TABLAS:
        raise HTTPException(status_code=404,
                            detail="Tabla desconocida. Validas: %s" % ", ".join(TABLAS))
    limit = max(1, min(int(limit), 500))
    offset = max(0, int(offset))

    conn = _conn()
    try:
        columnas = [r["name"] for r in conn.execute("PRAGMA table_info(%s)" % nombre)]
        where, params = "", []
        if q:
            # Busqueda simple sobre las columnas de texto.
            textos = [c for c in columnas if c.lower().endswith(("code", "class", "area", "group", "name"))]
            if textos:
                where = " WHERE " + " OR ".join("CAST(%s AS TEXT) LIKE ?" % c for c in textos)
                params = ["%%%s%%" % q] * len(textos)

        total = conn.execute("SELECT COUNT(*) FROM %s%s" % (nombre, where), params).fetchone()[0]
        filas = conn.execute(
            "SELECT * FROM %s%s LIMIT ? OFFSET ?" % (nombre, where), params + [limit, offset]
        ).fetchall()
        return {
            "tabla": nombre,
            "hoja": TABLAS[nombre],
            "columnas": columnas,
            "total": total,
            "limit": limit,
            "offset": offset,
            "filas": [dict(f) for f in filas],
        }
    finally:
        conn.close()
