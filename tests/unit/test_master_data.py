# -*- coding: utf-8 -*-
"""
Datos maestros gestionables desde la web.

El motor NO lee el Excel: lee `warehouse.db`. Editar el Excel no tenia efecto
hasta correr la migracion por consola, y la UI no lo avisaba. Estos tests
pinnean el contrato de los endpoints que resuelven eso.

Ver `docs/PLAN_DATOS_MAESTROS_WEB.md`.
"""
import os
import shutil

import openpyxl
import pytest

from web_prototype.app_state import PROJECT_ROOT
from web_prototype.routers import master_data as md

EXCEL_REAL = os.path.join(PROJECT_ROOT, "layouts", "Warehouse_Logic.xlsx")


@pytest.fixture()
def excel_copia(tmp_path):
    """Copia del Excel real: los tests no tocan el archivo del proyecto."""
    destino = tmp_path / "maestro.xlsx"
    shutil.copy2(EXCEL_REAL, destino)
    return str(destino)


# --------------------------------------------------------- validacion

def test_md01_excel_real_es_valido(excel_copia):
    r = md._validar_excel(excel_copia)
    assert r["valido"], r["errores"]
    assert r["resumen"]["PickingLocations"] > 0
    assert r["resumen"]["SkuCatalog"] > 0


def test_md02_falta_hoja_obligatoria(excel_copia, tmp_path):
    wb = openpyxl.load_workbook(excel_copia)
    del wb["PickingLocations"]
    roto = str(tmp_path / "sin_hoja.xlsx")
    wb.save(roto)

    r = md._validar_excel(roto)
    assert not r["valido"]
    assert any("PickingLocations" in e for e in r["errores"])


def test_md03_falta_una_columna(excel_copia, tmp_path):
    """Una hoja presente pero sin una columna requerida tambien se rechaza."""
    wb = openpyxl.load_workbook(excel_copia)
    ws = wb["PickingLocations"]
    for celda in ws[1]:
        if celda.value == "WorkArea":
            celda.value = "OtraCosa"
    roto = str(tmp_path / "sin_columna.xlsx")
    wb.save(roto)

    r = md._validar_excel(roto)
    assert not r["valido"]
    assert any("WorkArea" in e for e in r["errores"])


def test_md04_hoja_opcional_ausente_avisa_pero_no_rechaza(excel_copia, tmp_path):
    wb = openpyxl.load_workbook(excel_copia)
    del wb["SkuCatalog"]
    sin_opcional = str(tmp_path / "sin_opcional.xlsx")
    wb.save(sin_opcional)

    r = md._validar_excel(sin_opcional)
    assert r["valido"], r["errores"]
    assert any("SkuCatalog" in a for a in r["avisos"])


def test_md05_archivo_que_no_es_excel(tmp_path):
    falso = tmp_path / "no_es_excel.xlsx"
    falso.write_text("esto no es un xlsx", encoding="utf-8")
    r = md._validar_excel(str(falso))
    assert not r["valido"]
    assert r["errores"]


# ------------------------------------------------------------ rutas

def test_md06_bloquea_path_traversal():
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        md._ruta_segura(os.path.join("..", "..", "etc", "passwd"))


def test_md07_ruta_relativa_valida():
    assert md._ruta_segura("config.json").startswith(os.path.realpath(PROJECT_ROOT))


# ---------------------------------------------------------- resumen

def test_md08_resumen_reporta_las_tablas():
    r = md._resumen_bd()
    if not r.get("disponible"):
        pytest.skip("no hay warehouse.db en este entorno")
    for tabla in ("locations", "sku_catalog", "staging_areas", "inbound_docks"):
        assert tabla in r["conteos"]


def test_md09_detecta_excel_mas_nuevo(monkeypatch):
    """El aviso clave: el Excel cambio y todavia no se aplico.

    Sin esto, alguien edita el Excel, no ve ningun cambio en los resultados y
    no entiende por que.
    """
    if not os.path.exists(md.DB_PATH):
        pytest.skip("no hay warehouse.db en este entorno")

    real = os.path.getmtime

    def fake_mtime(path):
        # El Excel, 1000 s mas nuevo que la base.
        if str(path).lower().endswith(".xlsx"):
            return real(md.DB_PATH) + 1000
        return real(path)

    monkeypatch.setattr(os.path, "getmtime", fake_mtime)
    r = md.resumen()
    assert r.get("excel_mas_nuevo") is True


# ------------------------------------------- tablas chicas editables

def test_md10_solo_las_tablas_chicas_son_editables():
    """Las grandes se editan en el Excel, que es mejor herramienta para eso."""
    assert set(md.EDITABLES) == {"staging_areas", "inbound_docks"}
    assert "locations" not in md.EDITABLES
    assert "sku_catalog" not in md.EDITABLES


def test_md11_rechaza_ids_repetidos():
    from fastapi import HTTPException
    filas = [md.Coordenada(id=1, x=1, y=1), md.Coordenada(id=1, x=2, y=2)]
    with pytest.raises(HTTPException) as e:
        md._guardar_coordenadas("inbound_docks", filas)
    assert "repetidos" in str(e.value.detail).lower()


def test_md12_rechaza_coordenadas_negativas():
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as e:
        md._guardar_coordenadas("inbound_docks", [md.Coordenada(id=1, x=-5, y=1)])
    assert "negativ" in str(e.value.detail).lower()


def test_md13_rechaza_ids_inexistentes():
    """Agregar o quitar filas se hace en el Excel, no por esta via."""
    from fastapi import HTTPException
    if not os.path.exists(md.DB_PATH):
        pytest.skip("no hay warehouse.db en este entorno")
    with pytest.raises(HTTPException) as e:
        md._guardar_coordenadas("inbound_docks", [md.Coordenada(id=9999, x=1, y=1)])
    assert "no existen" in str(e.value.detail).lower()


def test_md15_rechaza_coordenadas_fuera_del_mapa():
    """El motor ABORTA la corrida si una zona/muelle cae fuera de la grilla del
    TMX. Sin este chequeo, la web aceptaba el valor y el error aparecia recien
    al simular, lejos de donde se cometio."""
    from fastapi import HTTPException
    if not os.path.exists(md.DB_PATH):
        pytest.skip("no hay warehouse.db en este entorno")
    ancho, alto, _ = md._mapa_para_validar()
    if ancho is None:
        pytest.skip("no se pudo cargar el mapa en este entorno")

    with pytest.raises(HTTPException) as e:
        md._guardar_coordenadas("inbound_docks",
                                [md.Coordenada(id=1, x=ancho + 50, y=alto + 50)])
    assert "fuera del mapa" in str(e.value.detail).lower()


def test_md16_rechaza_celdas_no_transitables():
    """Una celda puede estar dentro de la grilla y aun asi ser un rack: si se
    pone ahi un muelle, ningun agente puede llegar."""
    from fastapi import HTTPException
    if not os.path.exists(md.DB_PATH):
        pytest.skip("no hay warehouse.db en este entorno")
    ancho, alto, transitable = md._mapa_para_validar()
    if transitable is None:
        pytest.skip("no se pudo cargar el mapa en este entorno")

    bloqueada = next(((x, y) for y in range(alto) for x in range(ancho)
                      if not transitable(x, y)), None)
    if bloqueada is None:
        pytest.skip("el mapa no tiene celdas bloqueadas")

    with pytest.raises(HTTPException) as e:
        md._guardar_coordenadas("inbound_docks",
                                [md.Coordenada(id=1, x=bloqueada[0], y=bloqueada[1])])
    assert "transitable" in str(e.value.detail).lower()


def test_md14_guardar_y_restaurar_coordenadas():
    """Round-trip real contra la base, dejandola como estaba."""
    if not os.path.exists(md.DB_PATH):
        pytest.skip("no hay warehouse.db en este entorno")

    conn = md._conn()
    try:
        originales = [(r["dock_id"], r["x"], r["y"])
                      for r in conn.execute("SELECT dock_id, x, y FROM inbound_docks")]
    finally:
        conn.close()
    if not originales:
        pytest.skip("no hay muelles cargados")

    did, ox, oy = originales[0]

    # El destino tiene que ser una celda VALIDA (dentro del mapa y transitable),
    # porque el guardado ahora lo exige. Se busca una distinta de la actual.
    ancho, alto, transitable = md._mapa_para_validar()
    if ancho is None:
        destino = (ox + 1, oy)
    else:
        destino = next(((x, y) for y in range(alto) for x in range(ancho)
                        if (x, y) != (ox, oy)
                        and (transitable is None or transitable(x, y))), None)
        if destino is None:
            pytest.skip("no hay otra celda valida en el mapa")

    try:
        md._guardar_coordenadas("inbound_docks",
                                [md.Coordenada(id=did, x=destino[0], y=destino[1])])
        conn = md._conn()
        try:
            fila = conn.execute(
                "SELECT x, y FROM inbound_docks WHERE dock_id = ?", (did,)).fetchone()
        finally:
            conn.close()
        assert (fila["x"], fila["y"]) == destino
    finally:
        md._guardar_coordenadas("inbound_docks", [md.Coordenada(id=did, x=ox, y=oy)])

    conn = md._conn()
    try:
        fila = conn.execute(
            "SELECT x, y FROM inbound_docks WHERE dock_id = ?", (did,)).fetchone()
    finally:
        conn.close()
    assert (fila["x"], fila["y"]) == (ox, oy), "el test dejo la base modificada"
