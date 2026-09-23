# -*- coding: utf-8 -*-
"""
QA-6.6: al subir un mapa, la web lo valida CONTRA los datos maestros en uso.
Antes bastaba un <map>: se aceptaba un mapa vacio de 10x10 aunque las
ubicaciones del almacen llegan a x=31, y=39.
"""
import os
import sqlite3
import tempfile

import pytest

from web_prototype.routers.master_data import validar_tmx

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
V3 = os.path.join(RAIZ, 'layouts', 'WH1 v3.tmx')
V2 = os.path.join(RAIZ, 'layouts', 'WH1 v2.tmx')


def _archivo(texto):
    ruta = os.path.join(tempfile.mkdtemp(), 'm.tmx')
    open(ruta, 'w', encoding='utf-8').write(texto)
    return ruta


def _base(puntos):
    ruta = os.path.join(tempfile.mkdtemp(), 'wh.db')
    c = sqlite3.connect(ruta)
    c.execute("CREATE TABLE locations (location_id TEXT, legacy_x INT, legacy_y INT)")
    c.execute("CREATE TABLE staging_areas (staging_id INT, legacy_x INT, legacy_y INT)")
    c.execute("CREATE TABLE inbound_docks (dock_id INT, x INT, y INT)")
    c.executemany("INSERT INTO locations VALUES (?, ?, ?)", puntos)
    c.commit(); c.close()
    return ruta


def test_texto_que_no_es_xml_se_rechaza():
    r = validar_tmx(_archivo('esto no es un mapa'))
    assert not r['valido'] and 'XML' in r['errores'][0]


def test_mapa_sin_capas_se_rechaza():
    r = validar_tmx(_archivo('<?xml version="1.0"?><map width="10" height="10" '
                             'tilewidth="32" tileheight="32"></map>'))
    assert not r['valido']
    assert any('capas' in e for e in r['errores'])


@pytest.mark.skipif(not os.path.exists(V3), reason='sin WH1 v3')
def test_el_mapa_canonico_es_valido_con_sus_datos():
    r = validar_tmx(V3, _base([('LOC-001', 1, 3), ('LOC-384', 30, 3)]))
    assert r['valido'], r['errores']
    assert r['resumen']['celdas_bloqueadas'] > 0


@pytest.mark.skipif(not os.path.exists(V2), reason='sin WH1 v2')
def test_puntos_fuera_del_mapa_o_sobre_un_rack_se_rechazan():
    # WH1 v2 mide 30 de ancho: x=30 queda fuera. (0,3) es un rack.
    r = validar_tmx(V2, _base([('LOC-001', 1, 3), ('LOC-384', 30, 3), ('LOC-X', 0, 3)]))
    assert not r['valido']
    assert any('FUERA' in e and 'LOC-384' in e for e in r['errores'])
    assert any('bloqueadas' in e and 'LOC-X' in e for e in r['errores'])


# --- QA-6.8: el Excel tambien se cruza con el mapa ---------------------------
XLSX = os.path.join(RAIZ, 'layouts', 'Warehouse_Logic_v3.xlsx')


@pytest.mark.skipif(not (os.path.exists(XLSX) and os.path.exists(V3) and os.path.exists(V2)),
                    reason='sin Excel v3 o mapas')
def test_el_excel_se_valida_contra_el_mapa():
    from web_prototype.routers.master_data import _validar_excel
    assert _validar_excel(XLSX, V3)['valido']                       # canonico con su mapa
    r = _validar_excel(XLSX, V2)                                     # v2 es 2 columnas mas angosto
    assert not r['valido'] and any('FUERA' in e for e in r['errores'])


@pytest.mark.skipif(not (os.path.exists(XLSX) and os.path.exists(V3)), reason='sin Excel v3')
def test_un_carril_sobre_un_rack_se_rechaza():
    import openpyxl
    from web_prototype.routers.master_data import _validar_excel
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['OutboundStaging']
    ws.cell(row=2, column=2).value = 0          # zona 1, primera celda -> (0, 3): rack
    ws.cell(row=2, column=3).value = 3
    ruta = os.path.join(tempfile.mkdtemp(), 'x.xlsx')
    wb.save(ruta)
    r = _validar_excel(ruta, V3)
    assert not r['valido']
    assert any('bloqueadas' in e and 'zona de salida 1 (0, 3)' in e for e in r['errores'])
